"""
Actualización de los resultados oficiales del benchmark en el RESUMEN y en el ABSTRACT.
Incorpora las métricas oficiales de la Segunda Tanda de evaluación:
- Exactitud global fáctica: 45,00% (Sipán-STAIR) vs 6,00% (Gemma-4 Fine-Tuned)
- Tasa de alucinaciones: 48,0% (Sipán-STAIR) vs 90,0% (Gemma-4 Fine-Tuned)
- Latencia media: 7,47 s (Sipán-STAIR) vs 12,15 s (Gemma-4 Fine-Tuned)
Compatible de forma nativa con macOS y Windows.
"""

import sys
import re
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from platform_utils import (
    get_workspace_path,
    find_onedrive_path,
    create_backup,
    dual_save,
    evaluate_sync,
    DOCUMENT_DEFAULT_NAME,
    MARKDOWN_DEFAULT_NAME
)

try:
    import docx
    from docx.shared import Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
except ImportError:
    print("Error: python-docx no está instalado. Ejecute: pip install -r requirements.txt")
    sys.exit(1)

new_resumen = (
    "A partir del diagnóstico preliminar de operaciones de tecnologías de la información que reporta una demanda asistencial "
    "estimada de 1,400 a 1,500 solicitudes por ciclo académico con demoras de resolución de hasta 72 horas en la Universidad Señor de Sipán, "
    "la presente investigación tuvo como objetivo desarrollar un chatbot basado en arquitecturas Transformer para optimizar el soporte a las operaciones "
    "tecnológicas institucionales. La metodología correspondió a una investigación aplicada con enfoque mixto y diseño cuasi-experimental con post-prueba. "
    "La muestra documental estuvo conformada por trece manuales de usuario y reglamentos oficiales vigentes organizados en cinco módulos operativos. "
    "A partir de este acervo se sintetizó un corpus balanceado de 400 conversaciones en formato ShareGPT (70% monoturno y 30% multiturno), "
    "aplicándose ajuste fino mediante adaptación de bajo rango (LoRA) sobre el modelo razonador Gemma-4-E2B-it cuantizado a GGUF e implementándose "
    "paralelamente la arquitectura de generación aumentada por recuperación jerárquica Sipán-STAIR con indexación basada en árbol de contenidos sobre PostgreSQL. "
    "La evaluación técnica comparativa sobre un banco de pruebas ciego de 50 casos mediante la metodología de evaluación guiada por referencia "
    "(Reference-Guided LLM-as-a-Judge) evidenció que la arquitectura Sipán-STAIR alcanzó un 45,00% de exactitud global fáctica frente al 6,00% obtenido "
    "por el modelo con ajuste fino, logrando reducir la tasa de alucinaciones del 90,0% al 48,0%. Asimismo, la arquitectura Sipán-STAIR registró una latencia "
    "media de 7,47 segundos por consulta frente a 12,15 segundos de la inferencia local con el modelo adaptado, demostrando una mayor eficiencia temporal "
    "en el flujo de recuperación documental. La evaluación de usabilidad y satisfacción con una muestra de [Z] usuarios finales reportó una puntuación media "
    "de [X] puntos en la escala SUS y una satisfacción del [Y]%. Se concluye que la recuperación jerárquica sobre Transformers constituye una solución "
    "técnicamente viable y superior para garantizar rigor fáctico en el soporte tecnológico universitario."
)

new_abstract = (
    "Based on the preliminary information technology operations diagnosis reporting an estimated service demand of 1,400 to 1,500 support requests "
    "per academic semester with resolution delays of up to 72 hours at Universidad Señor de Sipán, this research aimed to develop an artificial intelligence "
    "chatbot based on Transformer architectures to optimize institutional technological operations support. The methodological design corresponded to applied "
    "research with a mixed approach and quasi-experimental post-test design. The documentary sample comprised thirteen active official user manuals and university "
    "regulations organized into five operational modules. From this institutional corpus, a balanced dataset of 400 conversations was synthesized in ShareGPT "
    "format (70% single-turn and 30% multi-turn), applying fine-tuning via Low-Rank Adaptation (LoRA) on the quantized Gemma-4-E2B-it reasoning model in GGUF format, "
    "and concurrently implementing the Sipán-STAIR hierarchical retrieval-augmented generation architecture with table-of-contents tree indexing on PostgreSQL. "
    "Technical comparative evaluation across an independent 50-case benchmark using the Reference-Guided LLM-as-a-Judge methodology revealed that the "
    "Sipán-STAIR architecture achieved 45.00% overall factual accuracy compared to 6.00% attained by the fine-tuned model, reducing the hallucination rate from "
    "90.0% to 48.0%. Furthermore, the Sipán-STAIR architecture recorded an average query latency of 7.47 seconds versus 12.15 seconds for local edge inference "
    "with the adapted model, demonstrating higher temporal efficiency in the document retrieval pipeline. Usability and satisfaction assessment conducted with a sample "
    "of [Z] end users reported an average score of [X] points on the System Usability Scale (SUS) and a user satisfaction level of [Y]%. It is concluded that "
    "structure-aware hierarchical retrieval over Transformers constitutes a technically viable, robust, and superior solution for ensuring factual correctness "
    "and operational efficiency in higher education technological support."
)

def update_docx(docx_file: Path) -> None:
    print(f"\n[BENCHMARK RESUMEN] Modificando {docx_file.name}...")
    doc = docx.Document(str(docx_file))

    for p in doc.paragraphs:
        t = p.text.strip()
        if "A partir del diagnóstico preliminar de operaciones de tecnologías de la información" in t:
            p.text = ""
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run(new_resumen)
            r.font.name = "Times New Roman"
            r.font.size = Pt(12)
            r.font.color.rgb = RGBColor(0, 0, 0)
            print("  RESUMEN actualizado.")

        elif "Based on the preliminary information technology operations diagnosis" in t:
            p.text = ""
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.line_spacing = 1.5
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run(new_abstract)
            r.font.name = "Times New Roman"
            r.font.size = Pt(12)
            r.font.color.rgb = RGBColor(0, 0, 0)
            print("  ABSTRACT actualizado.")

    ws_path, od_path = dual_save(doc, docx_file.name)
    if od_path:
        evaluate_sync(od_path)

def update_markdown(md_file: Path) -> None:
    print(f"\n[BENCHMARK RESUMEN] Modificando Markdown {md_file.name}...")
    with open(md_file, "r", encoding="utf-8") as f:
        content = f.read()

    resumen_pattern = r"(# RESUMEN\s*\n\n)A partir del diagnóstico preliminar.*?(?=\n\n\*\*\*Palabras claves:\*\*\*)"
    content = re.sub(resumen_pattern, rf"\g<1>{new_resumen}", content, flags=re.DOTALL)

    abstract_pattern = r"(# ABSTRACT\s*\n\n)Based on the preliminary information technology operations diagnosis.*?(?=\n\n\*\*\*Keywords:\*\*\*)"
    content = re.sub(abstract_pattern, rf"\g<1>{new_abstract}", content, flags=re.DOTALL)

    with open(md_file, "w", encoding="utf-8") as f:
        f.write(content)
    print("  Markdown sincronizado correctamente.")

def main():
    ws_docx = get_workspace_path(DOCUMENT_DEFAULT_NAME)
    md_file = get_workspace_path(MARKDOWN_DEFAULT_NAME)

    create_backup(ws_docx)
    update_docx(ws_docx)
    update_markdown(md_file)

if __name__ == "__main__":
    main()
